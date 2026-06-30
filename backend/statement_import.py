"""Broker / platform statement import.

Traders rarely have a clean column of returns -- they have an MT4/MT5 HTML
statement, a TradingView "List of Trades" export, or a cTrader CSV. This module
turns those multi-column exports into the single normalised numeric column the
existing parser + analysis pipeline already consumes. It does NOT analyse: it
only locates and extracts the right series (per-trade P&L, or a balance/equity
curve) and reports what it found, so the normal analyze flow can run unchanged.

Supported inputs:
  * MT4 / MT5 HTML statements (.htm / .html)        -> Profit column (per trade)
  * TradingView / MT5 XLSX reports (.xlsx)          -> P&L column (per trade)
  * cTrader / generic multi-column CSV / TSV        -> P&L or balance column
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser

MAX_BYTES = 5 * 1024 * 1024  # 5 MB, mirrors parsing.MAX_BYTES
MAX_ROWS = 100_000

# Header hints. Per-trade P&L is preferred over a running balance/equity column.
PNL_HINTS = ["profit", "p&l", "p/l", "pnl", "pl", "net", "result", "gain", "win/loss"]
EQUITY_HINTS = ["balance", "equity", "nav", "cumulative", "running"]
# Substrings that disqualify a column from being a *per-trade* P&L series.
PNL_EXCLUDE = ["%", "percent", "pct", "cumulative", "running", "total", "balance",
               "equity", "commission", "swap", "fee", "price", "pips", "points"]


class ImportError_(ValueError):
    """Raised when no usable series can be found in a statement."""


@dataclass
class ImportResult:
    values: list[float]
    kind: str               # "trades" | "equity"
    column: str             # the source column we read
    source_format: str      # "mt4_mt5_html" | "xlsx" | "csv"
    n: int = 0
    notes: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        """The extracted column as plain text the existing parser accepts."""
        header = "pnl" if self.kind == "trades" else "equity"
        return header + "\n" + "\n".join(_fmt(v) for v in self.values)

    def to_meta(self) -> dict:
        return {
            "source_format": self.source_format,
            "kind": self.kind,
            "column": self.column,
            "n": self.n,
            "notes": self.notes,
        }


def _fmt(v: float) -> str:
    return f"{v:.10g}"


# ---------------------------------------------------------------------------
# number parsing (broker exports use spaces / commas as thousands separators)
# ---------------------------------------------------------------------------

_NUM_RE = re.compile(r"-?\d[\d\s.,]*\d|-?\d")


def _parse_number(cell: str) -> float | None:
    if cell is None:
        return None
    s = str(cell).strip()
    if not s:
        return None
    # Keep a leading minus / parentheses-as-negative, drop currency & symbols.
    neg = s.startswith("(") and s.endswith(")")
    s = s.replace("(", "").replace(")", "")
    s = s.replace("\xa0", " ")  # non-breaking space (common in MT4 HTML)
    m = _NUM_RE.search(s)
    if not m:
        return None
    tok = m.group(0).strip()
    has_dot = "." in tok
    has_comma = "," in tok
    if has_dot and has_comma:
        # Both present: comma is the thousands separator (e.g. 1,234.56).
        tok = tok.replace(",", "")
    elif has_comma and not has_dot:
        # Comma only: decimal comma if it looks like one (1,5), else thousands.
        if re.search(r",\d{1,2}$", tok) and tok.count(",") == 1:
            tok = tok.replace(",", ".")
        else:
            tok = tok.replace(",", "")
    tok = tok.replace(" ", "")
    try:
        val = float(tok)
    except ValueError:
        return None
    return -val if neg else val


def _is_number(cell: str) -> bool:
    return _parse_number(cell) is not None


# ---------------------------------------------------------------------------
# column selection over a generic table (list of rows of string cells)
# ---------------------------------------------------------------------------

def _header_score(name: str, hints: list[str]) -> bool:
    n = name.strip().lower()
    return any(h in n for h in hints)


def _select_column(rows: list[list[str]]) -> tuple[list[float], str, str, list[str]]:
    """Pick the best P&L (preferred) or equity column from a table.

    Returns (values, kind, column_label, notes).
    """
    notes: list[str] = []
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ImportError_("The statement had no readable rows.")

    # Find the header row: the first row that is mostly non-numeric text.
    header_idx = 0
    for i, r in enumerate(rows[:5]):
        non_numeric = sum(1 for c in r if c and not _is_number(c))
        if non_numeric >= max(2, len(r) // 2):
            header_idx = i
            break
    header = [str(c).strip() for c in rows[header_idx]]
    body = rows[header_idx + 1:]
    ncols = max((len(r) for r in body), default=len(header))

    # Build per-column numeric values aligned to header positions.
    def col_values(j: int) -> list[float]:
        out = []
        for r in body:
            if j < len(r):
                v = _parse_number(r[j])
                if v is not None:
                    out.append(v)
        return out

    labelled = []  # (j, label)
    for j in range(ncols):
        label = header[j] if j < len(header) else f"column {j + 1}"
        labelled.append((j, label or f"column {j + 1}"))

    # 1) prefer a per-trade P&L column by header hint.
    for j, label in labelled:
        low = label.lower()
        if _header_score(low, PNL_HINTS) and not any(x in low for x in PNL_EXCLUDE):
            vals = col_values(j)
            if len(vals) >= 2:
                return vals, "trades", label, notes

    # 2) else an equity / balance column by header hint.
    for j, label in labelled:
        if _header_score(label.lower(), EQUITY_HINTS):
            vals = col_values(j)
            if len(vals) >= 2:
                notes.append("No per-trade P&L column found; used the "
                             f"'{label}' balance/equity column instead.")
                return vals, "equity", label, notes

    # 3) fallback: the numeric column with the most values that isn't an index.
    best = None
    for j, label in labelled:
        low = label.lower()
        if low in ("#", "no", "no.", "ticket", "id", "index", "trade #", "order"):
            continue
        vals = col_values(j)
        if best is None or len(vals) > len(best[0]):
            best = (vals, label)
    if best and len(best[0]) >= 2:
        notes.append(f"Could not identify a P&L column by name; analysed the "
                     f"most-populated numeric column ('{best[1]}').")
        return best[0], "trades", best[1], notes

    raise ImportError_("Could not find a profit/P&L or balance column to read.")


# ---------------------------------------------------------------------------
# format parsers
# ---------------------------------------------------------------------------

class _TableHTMLParser(HTMLParser):
    """Collects every <table> as a list of rows, each a list of cell text."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._table = []
            self.tables.append(self._table)
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._table is not None:
                self._table.append(self._row)
            self._row = None
        elif tag == "table":
            self._table = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _parse_html(content: str) -> ImportResult:
    p = _TableHTMLParser()
    p.feed(content)
    # Pick the table that yields the best (largest) P&L/equity series.
    best: ImportResult | None = None
    for table in p.tables:
        if len(table) < 2:
            continue
        try:
            vals, kind, col, notes = _select_column(table)
        except ImportError_:
            continue
        cand = ImportResult(values=vals, kind=kind, column=col,
                            source_format="mt4_mt5_html", n=len(vals), notes=notes)
        if best is None or cand.n > best.n:
            best = cand
    if best is None:
        raise ImportError_("No usable table found in the HTML statement.")
    return best


def _parse_xlsx(content: bytes) -> ImportResult:
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - dependency is declared
        raise ImportError_("XLSX support isn't available on this server.")
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    best: ImportResult | None = None
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                break
            rows.append(["" if c is None else str(c) for c in row])
        if len(rows) < 2:
            continue
        try:
            vals, kind, col, notes = _select_column(rows)
        except ImportError_:
            continue
        if len(wb.worksheets) > 1:
            notes = notes + [f"Read sheet '{ws.title}'."]
        cand = ImportResult(values=vals, kind=kind, column=col,
                            source_format="xlsx", n=len(vals), notes=notes)
        if best is None or cand.n > best.n:
            best = cand
    wb.close()
    if best is None:
        raise ImportError_("No usable sheet found in the workbook.")
    return best


def _parse_csv(content: str) -> ImportResult:
    import csv
    sample = content[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    rows = []
    for i, r in enumerate(csv.reader(io.StringIO(content), delimiter=delim)):
        if i >= MAX_ROWS:
            break
        rows.append([c.strip() for c in r])
    vals, kind, col, notes = _select_column(rows)
    return ImportResult(values=vals, kind=kind, column=col,
                        source_format="csv", n=len(vals), notes=notes)


# ---------------------------------------------------------------------------
# public entry point
# ---------------------------------------------------------------------------

def parse_statement(content: bytes, filename: str = "") -> ImportResult:
    """Detect the export format and extract its P&L / equity series."""
    if content is None or len(content) == 0:
        raise ImportError_("The uploaded file was empty.")
    if len(content) > MAX_BYTES:
        raise ImportError_("File exceeds the 5 MB limit.")

    name = (filename or "").lower()
    # XLSX is binary (zip): sniff the PK header or the extension.
    if name.endswith(".xlsx") or content[:2] == b"PK":
        return _parse_xlsx(content)

    text = content.decode("utf-8", errors="replace")
    stripped = text.lstrip().lower()
    if name.endswith((".htm", ".html")) or stripped.startswith(("<!doctype", "<html")) or "<table" in stripped:
        return _parse_html(text)

    return _parse_csv(text)
